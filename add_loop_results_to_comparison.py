#!/usr/bin/env python3
"""Add selection-loop results to the existing method comparison workbook.

Example:
  python add_loop_results_to_comparison.py \
    --comparison-xlsx ../method1_method2_claude_all_results.xlsx \
    --loop-results /tmp/oren15_all75_full5000_final_results.csv \
    --xlsx-output ../method1_method2_claude_loop_all_results.xlsx \
    --csv-output ../method1_method2_claude_loop_all_results.csv
"""
from __future__ import annotations

import argparse
import csv
import math
from collections import defaultdict
from pathlib import Path


SUCCESS_STATUSES = {"complete", "ok", "success", "succeeded", "warn"}
LOOP_FIELDS = ("loop_model", "loop_score", "loop_metric", "loop_count")
SCORE_FIELDS = {
    "method1_score",
    "method2_score",
    "claude_score",
    "loop_score",
}
COUNT_FIELDS = {
    "method1_count",
    "method2_count",
    "claude_count",
    "loop_count",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--comparison-xlsx", type=Path, required=True)
    parser.add_argument("--loop-results", type=Path, required=True)
    parser.add_argument("--xlsx-output", type=Path, required=True)
    parser.add_argument("--csv-output", type=Path, required=True)
    return parser.parse_args()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def numeric_value(value: object) -> float | None:
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def is_measurement(row: dict[str, str]) -> bool:
    status = row.get("status", "").lower()
    score = numeric_value(row.get("score"))
    labeled = numeric_value(row.get("labeled_total"))
    return status in SUCCESS_STATUSES and score is not None and (labeled or 0) > 0


def selection_rank(row: dict[str, str]) -> int:
    value = row.get("selection_rank") or row.get("final_rank") or 0
    return int(value)


def eligible_rows(rows: list[dict[str, str]], max_rank: int) -> list[dict[str, str]]:
    return [
        row
        for row in rows
        if selection_rank(row) <= max_rank and is_measurement(row)
    ]


def score_text(row: dict[str, str] | None) -> str:
    score = numeric_value(row.get("score")) if row else None
    return f"{score:.9g}" if score is not None else ""


def loop_pool(rows: list[dict[str, str]], max_rank: int) -> dict[str, str]:
    measured = eligible_rows(rows, max_rank)
    best = max(measured, key=lambda row: float(row["score"]), default=None)
    return {
        "loop_model": best.get("model", "") if best else "",
        "loop_score": score_text(best),
        "loop_metric": best.get("metric", "") if best else "",
        "loop_count": f"{len(measured)}/{max_rank}",
    }


def build_loop_results(rows: list[dict[str, str]]) -> dict[tuple[str, str], dict[str, str]]:
    grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        grouped[row["dataset"]].append(row)
    output = {}
    for dataset, dataset_rows in grouped.items():
        output[(dataset, "top1")] = loop_pool(dataset_rows, 1)
        output[(dataset, "best_of_top5")] = loop_pool(dataset_rows, 5)
    return output


def selected_groups(row: dict[str, object]) -> list[tuple[str, float, str]]:
    groups = []
    for name in ("method1", "method2", "claude", "loop"):
        score = numeric_value(row.get(f"{name}_score"))
        model = row.get(f"{name}_model")
        if score is not None and model:
            groups.append((name, score, str(row.get(f"{name}_metric") or "")))
    return groups


def winner_with_loop(row: dict[str, object]) -> str:
    groups = selected_groups(row)
    if len(groups) != 4:
        return "incomplete"
    metrics = {metric for _, _, metric in groups if metric}
    if len(metrics) > 1:
        return "incomplete"
    maximum = max(score for _, score, _ in groups)
    winners = [
        name
        for name, score, _ in groups
        if math.isclose(score, maximum, rel_tol=1e-9, abs_tol=1e-12)
    ]
    return winners[0] if len(winners) == 1 else "tie"


def augment_row(
    row: dict[str, object], loop_results: dict[tuple[str, str], dict[str, str]]
) -> dict[str, object]:
    loop = loop_results.get((str(row["dataset"]), str(row["comparison"])))
    if not loop:
        return {**row, **dict.fromkeys(LOOP_FIELDS, "")}
    augmented = {**row, **loop}
    augmented["winner"] = winner_with_loop(augmented)
    return augmented


def ordered_fields(fields: list[str]) -> list[str]:
    without_winner = [field for field in fields if field != "winner"]
    return [*without_winner, *LOOP_FIELDS, "winner"]


def read_xlsx(path: Path) -> tuple[list[str], list[dict[str, object]]]:
    from openpyxl import load_workbook

    worksheet = load_workbook(path, read_only=True, data_only=True).active
    values = worksheet.iter_rows(values_only=True)
    fields = [str(value) for value in next(values)]
    rows = [dict(zip(fields, values_row)) for values_row in values]
    return fields, rows


def output_rows(
    fields: list[str],
    rows: list[dict[str, object]],
    loop_results: dict[tuple[str, str], dict[str, str]],
) -> tuple[list[str], list[dict[str, object]]]:
    output_fields = ordered_fields(fields)
    return output_fields, [augment_row(row, loop_results) for row in rows]


def write_csv(path: Path, fields: list[str], rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(
            {field: csv_value(field, row.get(field, "")) for field in fields}
            for row in rows
        )


def csv_value(field: str, value: object) -> object:
    if field in COUNT_FIELDS and value not in (None, ""):
        return f"'{value}"
    return "" if value is None else value


def xlsx_value(field: str, value: object) -> object:
    if field in SCORE_FIELDS and value not in (None, ""):
        return float(value)
    return "" if value is None else value


def set_column_widths(worksheet, fields: list[str], rows: list[dict[str, object]]) -> None:
    from openpyxl.utils import get_column_letter

    for index, field in enumerate(fields, start=1):
        width = max(len(field), *(len(str(row.get(field, ""))) for row in rows)) + 2
        worksheet.column_dimensions[get_column_letter(index)].width = min(width, 60)


def write_xlsx(
    path: Path, fields: list[str], rows: list[dict[str, object]]
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "comparison"
    worksheet.append(fields)
    for row in rows:
        worksheet.append([xlsx_value(field, row.get(field, "")) for field in fields])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for index, field in enumerate(fields, start=1):
        if field in COUNT_FIELDS:
            for cell in worksheet.iter_cols(
                min_col=index, max_col=index, min_row=2, max_row=worksheet.max_row
            ):
                for value in cell:
                    value.number_format = "@"
    set_column_widths(worksheet, fields, rows)
    workbook.save(path)


def main() -> int:
    args = parse_args()
    fields, existing = read_xlsx(args.comparison_xlsx)
    loop_results = build_loop_results(read_csv(args.loop_results))
    output_fields, rows = output_rows(fields, existing, loop_results)
    write_csv(args.csv_output, output_fields, rows)
    write_xlsx(args.xlsx_output, output_fields, rows)
    matched = sum(bool(row.get("loop_model")) for row in rows)
    print(f"Added loop results to {matched}/{len(rows)} comparison rows.")
    print(f"CSV: {args.csv_output}")
    print(f"XLSX: {args.xlsx_output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
