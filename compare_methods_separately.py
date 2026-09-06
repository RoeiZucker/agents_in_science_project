#!/usr/bin/env python3
"""Compare method1, method2, and Claude as separate candidate pools.

Example:
  python compare_methods_separately.py \
    --input /path/to/all_results_pairs.csv \
    --csv-output /path/to/method1_method2_claude_all_results.csv \
    --markdown-output /path/to/method1_method2_claude_all_results.md \
    --xlsx-output /path/to/method1_method2_claude_all_results.xlsx
"""
from __future__ import annotations

import argparse
import csv
import math
from collections import Counter
from pathlib import Path

from compare_methods_to_claude import (
    best_candidate,
    candidate_pool,
    coverage,
    numeric_score,
    read_rows,
    score_text,
)


METHOD1 = {"B_method1_only"}
METHOD2 = {"C_method2_only"}
CLAUDE = {"D_claude_baseline"}
EXCEL_TEXT_FIELDS = {
    "method1_count",
    "method2_count",
    "claude_count",
}
SCORE_FIELDS = {"method1_score", "method2_score", "claude_score"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--csv-output", required=True, type=Path)
    parser.add_argument("--markdown-output", required=True, type=Path)
    parser.add_argument("--xlsx-output", type=Path)
    return parser.parse_args()


def known_metrics(rows: list[dict[str, str] | None]) -> set[str]:
    return {row.get("eval_metric", "") for row in rows if row and row.get("eval_metric")}


def three_way_winner(
    method1: dict[str, str] | None,
    method2: dict[str, str] | None,
    claude: dict[str, str] | None,
) -> str:
    candidates = [method1, method2, claude]
    if any(candidate is None for candidate in candidates):
        return "incomplete"
    if len(known_metrics(candidates)) > 1:
        return "incomplete"
    scores = [numeric_score(candidate) for candidate in candidates if candidate]
    if any(score is None for score in scores):
        return "incomplete"
    maximum = max(score for score in scores if score is not None)
    winners = [
        name
        for name, score in zip(("method1", "method2", "claude"), scores)
        if score is not None and math.isclose(score, maximum, rel_tol=1e-9, abs_tol=1e-12)
    ]
    return winners[0] if len(winners) == 1 else "tie"


def pool_result(
    rows: list[dict[str, str]], names: set[str], max_rank: int
) -> tuple[dict[str, str] | None, str]:
    pool = candidate_pool(rows, names, max_rank)
    return best_candidate(pool), coverage(pool)


def comparison(dataset: str, rows: list[dict[str, str]], max_rank: int) -> dict[str, str]:
    method1, method1_count = pool_result(rows, METHOD1, max_rank)
    method2, method2_count = pool_result(rows, METHOD2, max_rank)
    claude, claude_count = pool_result(rows, CLAUDE, max_rank)
    return {
        "dataset": dataset,
        "comparison": "top1" if max_rank == 1 else "best_of_top5",
        "method1_model": method1.get("model", "") if method1 else "",
        "method1_score": score_text(method1),
        "method1_metric": method1.get("eval_metric", "") if method1 else "",
        "method1_count": method1_count,
        "method2_model": method2.get("model", "") if method2 else "",
        "method2_score": score_text(method2),
        "method2_metric": method2.get("eval_metric", "") if method2 else "",
        "method2_count": method2_count,
        "claude_model": claude.get("model", "") if claude else "",
        "claude_score": score_text(claude),
        "claude_metric": claude.get("eval_metric", "") if claude else "",
        "claude_count": claude_count,
        "winner": three_way_winner(method1, method2, claude),
    }


def build_comparisons(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    datasets = list(dict.fromkeys(row.get("dataset", "") for row in rows))
    output = []
    for dataset in datasets:
        dataset_rows = [row for row in rows if row.get("dataset") == dataset]
        output.append(comparison(dataset, dataset_rows, 1))
        output.append(comparison(dataset, dataset_rows, 5))
    return output


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(excel_safe_row(row) for row in rows)


def excel_safe_row(row: dict[str, str]) -> dict[str, str]:
    return {
        name: f"'{value}" if name in EXCEL_TEXT_FIELDS and value else value
        for name, value in row.items()
    }


def xlsx_value(name: str, value: str) -> str | float:
    if name in SCORE_FIELDS and value:
        return float(value)
    return value


def set_column_widths(worksheet, rows: list[dict[str, str]]) -> None:
    from openpyxl.utils import get_column_letter

    for index, name in enumerate(rows[0], start=1):
        width = max(len(name), *(len(str(row[name])) for row in rows)) + 2
        worksheet.column_dimensions[get_column_letter(index)].width = min(width, 60)


def format_xlsx_sheet(worksheet, rows: list[dict[str, str]]) -> None:
    from openpyxl.styles import Font, PatternFill

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for column, name in enumerate(rows[0], start=1):
        if name in EXCEL_TEXT_FIELDS:
            for cells in worksheet.iter_cols(
                min_col=column, max_col=column, min_row=2
            ):
                for cell in cells:
                    cell.number_format = "@"
    set_column_widths(worksheet, rows)


def write_xlsx(path: Path, rows: list[dict[str, str]]) -> None:
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "comparison"
    fields = list(rows[0])
    worksheet.append(fields)
    for row in rows:
        worksheet.append([xlsx_value(name, row[name]) for name in fields])
    format_xlsx_sheet(worksheet, rows)
    workbook.save(path)


def write_markdown(path: Path, rows: list[dict[str, str]]) -> None:
    counts = Counter(row["winner"] for row in rows)
    lines = [
        "# Method1 vs Method2 vs Claude",
        "",
        "Counts are measured/eligible candidates. A winner is incomplete when any pool lacks a score or selected known metrics differ.",
        "",
        "| Dataset | Comparison | Method1: model, score, metric, count | Method2: model, score, metric, count | Claude: model, score, metric, count | Winner |",
        "|---|---|---|---|---|---|",
    ]
    for row in rows:
        values = [
            row["dataset"],
            row["comparison"],
            f"{row['method1_model']}; {row['method1_score']}; {row['method1_metric']}; {row['method1_count']}",
            f"{row['method2_model']}; {row['method2_score']}; {row['method2_metric']}; {row['method2_count']}",
            f"{row['claude_model']}; {row['claude_score']}; {row['claude_metric']}; {row['claude_count']}",
            row["winner"],
        ]
        lines.append("| " + " | ".join(value.replace("|", "\\|") for value in values) + " |")
    lines.extend(["", "## Outcomes", ""])
    lines.append(", ".join(f"{name}={count}" for name, count in sorted(counts.items())) + ".")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    args = parse_args()
    rows = build_comparisons(read_rows(args.input))
    write_csv(args.csv_output, rows)
    write_markdown(args.markdown_output, rows)
    if args.xlsx_output:
        write_xlsx(args.xlsx_output, rows)
    print(f"Wrote {len(rows)} separate-pool comparisons.")


if __name__ == "__main__":
    main()
